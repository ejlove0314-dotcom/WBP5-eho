#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_fig1.py — Figure 1 (Computational workflow) 재작성  2026-10-22
==================================================================
원고 Fig. 1 범례(개정판, 본문 문단 18)와 1:1로 대응하도록 다시 그린다.
모든 요소를 PowerPoint 네이티브 도형으로 만들므로 PPT에서 바로 편집할 수 있다.

이전 판 대비 수정
  1. Phase I 결론  "not enriched in malignant cells" (오류) ->
                   "not malignant-cell-specific" + lymphoid 부재 문구 추가
  2. MTiOpenScreen library-scale calibration 상자 신설 (누락분)
  3. "Exploratory structural ligandability" -> structural and chemical
  4. "do NOT validate" -> "do not validate"
  5. 붙임표 -> 엔대시 (29-104 -> 29–104, Phase II-III -> Phases II–III)
  6. "Credibility suite" -> "Robustness checks"  (cross-cohort 중복 제거)
  7. P0 에 잔기 범위(38–58) 추가

출력: Figure_1_workflow_20261022.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

C = lambda h: RGBColor.from_string(h)

BLUE_FILL, BLUE_LINE, BLUE_TEXT = C("EAF3FB"), C("2E75B6"), C("1F4E79")
PUR_HEAD, PUR_LINE, PUR_FILL, PUR_TEXT = C("7030A0"), C("7030A0"), C("E9E1F4"), C("4E2583")
PUR_SUBLN = C("B9A5D9")
GRN_HEAD, GRN_LINE, GRN_FILL, GRN_TEXT = C("2E7D32"), C("2E7D32"), C("E4F0DC"), C("2F5B22")
GRN_SUBLN = C("A9D08E")
WHITE = C("FFFFFF")

FONT = "Calibri"


def add(slide, shape, x, y, w, h, fill, line, lw=1.0, adj=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    if adj is not None and s.adjustments:
        s.adjustments[0] = adj
    s.text_frame.word_wrap = True
    return s


def text(shape, lines, size=10, color=None, bold=False, italic=False,
         align=PP_ALIGN.CENTER, space=0):
    """lines: str 또는 (text, bold, italic, size) 튜플의 리스트."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        if isinstance(ln, str):
            ln = (ln, bold, italic, size)
        t, b, it, sz = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = t
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(sz), b, it
        f.color.rgb = color if color else C("000000")
    return shape


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])

    L, W = 0.35, 12.63           # 좌측 여백, 전체 폭
    R = L + W

    # ── 1. 상단 전제 상자 ────────────────────────────────────────────
    b = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L, 0.22, W, 0.52,
            BLUE_FILL, BLUE_LINE, 1.25, adj=0.22)
    text(b, "WBP5 (TCEAL9): prognostic candidate carried forward from our prior "
            "bulk-tissue study [11] — not an unbiased discovery screen",
         size=12, color=BLUE_TEXT, bold=True)

    # ── 2. Phase I 컨테이너 ──────────────────────────────────────────
    p1_y, p1_h = 0.95, 3.56
    add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L, p1_y, W, p1_h, None, PUR_LINE, 1.5, adj=0.05)

    hd = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L + 0.08, p1_y + 0.07, W - 0.16, 0.42,
             PUR_HEAD, PUR_HEAD, 1.0, adj=0.30)
    text(hd, "Phase I — Transcriptomic target nomination  (main text)",
         size=12.5, color=WHITE, bold=True)

    cw, gap = 5.97, 0.25         # 2열 폭 / 간격
    cx1, cx2 = L + 0.25, L + 0.25 + cw + gap

    r1y, r1h = 1.58, 0.74
    s = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx1, r1y, cw, r1h, PUR_FILL, PUR_SUBLN, 1.0, adj=0.14)
    text(s, [("Discovery cohort: GSE103322 (HNSCC; 5,902 cells)", True, False, 10.5),
             ("Exploratory cohort: GSE198315 (HPV-negative OSCC)", True, False, 10.5)],
         color=PUR_TEXT, space=2)

    s = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx2, r1y, cw, r1h, PUR_FILL, PUR_SUBLN, 1.0, adj=0.14)
    text(s, [("Within-patient co-expression", True, False, 10.5),
             ("+ compartment-resolved analysis", True, False, 10.5)],
         color=PUR_TEXT, space=2)

    r2y, r2h = 2.42, 0.82
    s = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx1, r2y, cw, r2h, PUR_FILL, PUR_SUBLN, 1.0, adj=0.13)
    text(s, [("Program-level p-EMT enrichment", True, False, 10.5),
             ("depth-adjusted; expression-matched null gene sets", False, False, 9.5)],
         color=PUR_TEXT, space=2)

    s = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx2, r2y, cw, r2h, PUR_FILL, PUR_SUBLN, 1.0, adj=0.13)
    text(s, [("Robustness checks", True, False, 10.5),
             ("leave-one-patient-out; gene-family specificity; Xq22 module "
              "adjustment; negative and positive controls", False, False, 9.5)],
         color=PUR_TEXT, space=2)

    tw = 6.55
    s = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L + (W - tw) / 2, 3.33, tw, 0.46,
            PUR_FILL, PUR_SUBLN, 1.0, adj=0.22)
    text(s, "TCGA HNSC bulk cohort (n = 520; purity-adjusted) — complementary",
         size=10.5, color=PUR_TEXT, bold=True)

    s = add(sl, MSO_SHAPE.RECTANGLE, L + 0.25, 3.86, W - 0.50, 0.56, None, None)
    text(s, ["Finding: WBP5 marks a p-EMT program shared by the malignant and "
             "mesenchymal compartments and is largely absent from lymphoid cells;",
             "expression is not malignant-cell-specific, and the association is "
             "program-level rather than marker-level"],
         size=10, color=PUR_TEXT, italic=True, space=1)

    # ── 3. Phase II–III 컨테이너 ────────────────────────────────────
    p2_y, p2_h = 4.70, 2.58
    add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L, p2_y, W, p2_h, None, GRN_LINE, 1.5, adj=0.07)

    hd = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L + 0.08, p2_y + 0.07, W - 0.16, 0.42,
             GRN_HEAD, GRN_HEAD, 1.0, adj=0.30)
    text(hd, "Phases II–III — Exploratory structural and chemical "
             "ligandability  (Supplementary Information)",
         size=12.5, color=WHITE, bold=True)

    boxes = [
        [("AlphaFold2 model", True, False, 10.5),
         ("AF-Q9UHQ7-F1-model_v6", False, False, 9.0),
         ("76-residue core, 29–104", False, False, 9.0)],
        [("DoGSiteScorer", True, False, 10.5),
         ("candidate pocket P₀", False, False, 9.0),
         ("residues 38–58; Drug Score 0.84", False, False, 9.0)],
        [("AutoDock Vina", True, False, 10.5),
         ("docking of approved", False, False, 9.0),
         ("compounds", False, False, 9.0)],
        [("MTiOpenScreen", True, False, 10.5),
         ("library-scale calibration", False, False, 9.0),
         ("(Drugs-lib)", False, False, 9.0)],
        [("REINVENT4", True, False, 10.5),
         ("reinforcement-learning", False, False, 9.0),
         ("de novo generation", False, False, 9.0)],
        [("SwissADME / ProTox", True, False, 10.5),
         ("drug-likeness and", False, False, 9.0),
         ("predicted toxicity", False, False, 9.0)],
    ]
    n = len(boxes)
    g = 0.15
    bw = (W - 0.50 - g * (n - 1)) / n
    by, bh = 5.32, 1.12
    for i, lines in enumerate(boxes):
        s = add(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L + 0.25 + i * (bw + g), by, bw, bh,
                GRN_FILL, GRN_SUBLN, 1.0, adj=0.10)
        text(s, lines, color=GRN_TEXT, space=1)

    s = add(sl, MSO_SHAPE.RECTANGLE, L + 0.25, 6.56, W - 0.50, 0.55, None, None)
    text(s, ["Computationally prioritized hypotheses: these results do not validate "
             "the transcriptomic target nomination,",
             "and experimental validation is required"],
         size=10, color=GRN_TEXT, italic=True, space=1)

    # ── 4. 단계 사이 화살표 ─────────────────────────────────────────
    for y0, y1 in ((0.74, 0.95), (4.51, 4.70)):
        a = add(sl, MSO_SHAPE.DOWN_ARROW, L + W / 2 - 0.11, y0, 0.22, y1 - y0,
                C("A6A6A6"), None)
        a.line.fill.background()

    out = "Figure_1_workflow_20261022.pptx"
    prs.save(out)
    print("saved ->", out)


if __name__ == "__main__":
    main()
